import os, sqlite3, io, re
from datetime import date, datetime
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, abort
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

app=Flask(__name__); app.secret_key=os.getenv('SECRET_KEY','change-this-secret-key')
DB=os.getenv('DATABASE_PATH',os.path.join(os.path.dirname(__file__),'data.db'))
SAMPLE={'name':'新民國小跑道整修工程','office':'徐英豪建築師事務所','year':115,'prefix':'一一五所新環字第','recipient':'彰化縣田中鎮新民國民小學','speed':'普通件','security':'普通','address':'臺中市西屯區西屯路三段 159-72 號 5 樓','phone':'04-24631919','fax':'04-24631509','email':'yh6726.mail@msa.hinet.net','contact':'黃先生'}

def conn():
 c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; c.execute('PRAGMA foreign_keys=ON'); return c

def init_db():
 c=conn(); c.executescript('''
 CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,office TEXT NOT NULL,year INTEGER NOT NULL,prefix TEXT NOT NULL,default_recipient TEXT,speed TEXT,security TEXT,address TEXT,phone TEXT,fax TEXT,email TEXT,contact TEXT,created_at TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS documents(id INTEGER PRIMARY KEY AUTOINCREMENT,project_id INTEGER NOT NULL,doc_date TEXT NOT NULL,serial TEXT NOT NULL,recipient TEXT,speed TEXT,security TEXT,attachment TEXT,subject TEXT,explanation TEXT,originals TEXT,copies TEXT,created_at TEXT NOT NULL,updated_at TEXT NOT NULL,FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE);
 CREATE TABLE IF NOT EXISTS templates(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,subject TEXT,explanation TEXT,attachment TEXT,originals TEXT,copies TEXT,quick_enabled INTEGER DEFAULT 0,quick_buttons TEXT,created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
 ''')
 if c.execute('SELECT COUNT(*) FROM projects').fetchone()[0]==0:
  c.execute('INSERT INTO projects(name,office,year,prefix,default_recipient,speed,security,address,phone,fax,email,contact,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(SAMPLE['name'],SAMPLE['office'],SAMPLE['year'],SAMPLE['prefix'],SAMPLE['recipient'],SAMPLE['speed'],SAMPLE['security'],SAMPLE['address'],SAMPLE['phone'],SAMPLE['fax'],SAMPLE['email'],SAMPLE['contact'],datetime.now().isoformat(timespec='seconds')))
 if c.execute('SELECT COUNT(*) FROM templates').fetchone()[0]==0:
  now=datetime.now().isoformat(timespec='seconds'); c.execute('INSERT INTO templates(name,subject,explanation,attachment,originals,copies,quick_enabled,quick_buttons,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)',('材料送審','本所辦理「{{工程名稱}}」，檢送本案材料送審資料，詳如說明，敬請 鑒核。','依契約規定辦理。\n{{快速內容}}','','{{受文者}}','',1,'廠商證明文件,規格樣式,試驗報告,防火證明,型錄,出廠證明',now,now))
 c.commit(); c.close()

def get_project(pid):
 c=conn(); r=c.execute('SELECT * FROM projects WHERE id=?',(pid,)).fetchone(); c.close(); return r
def get_template(tid):
 c=conn(); r=c.execute('SELECT * FROM templates WHERE id=?',(tid,)).fetchone(); c.close(); return r
def all_projects():
 c=conn(); r=c.execute('SELECT * FROM projects ORDER BY id DESC').fetchall(); c.close(); return r
def all_templates():
 c=conn(); r=c.execute('SELECT * FROM templates ORDER BY id DESC').fetchall(); c.close(); return r
def project_docs(pid):
 c=conn(); r=c.execute('SELECT * FROM documents WHERE project_id=? ORDER BY id DESC',(pid,)).fetchall(); c.close(); return r
def all_docs(q=''):
 c=conn()
 if q:r=c.execute('SELECT d.*,p.name project_name FROM documents d JOIN projects p ON p.id=d.project_id WHERE d.subject LIKE ? OR d.recipient LIKE ? OR p.name LIKE ? ORDER BY d.id DESC',(f'%{q}%',f'%{q}%',f'%{q}%')).fetchall()
 else:r=c.execute('SELECT d.*,p.name project_name FROM documents d JOIN projects p ON p.id=d.project_id ORDER BY d.id DESC').fetchall()
 c.close(); return r

def next_serial(pid,d):
 c=conn(); n=c.execute('SELECT COUNT(*) FROM documents WHERE project_id=? AND doc_date=?',(pid,d)).fetchone()[0]+1; c.close(); return f'{n:02d}'
def doc_no(p,d,serial):
 dt=datetime.strptime(d,'%Y-%m-%d'); return f'{p["prefix"]} {dt.year-1911:03d}{dt.month:02d}{dt.day:02d}-{serial} 號'
def substitute(text,p,extra=None):
 text=text or ''; dt=date.today(); vals={'工程名稱':p['name'],'專案名稱':p['name'],'事務所':p['office'],'受文者':p['default_recipient'] or '','發文日期':f'中華民國 {dt.year-1911} 年 {dt.month:02d} 月 {dt.day:02d} 日','民國年':str(dt.year-1911),'月份':f'{dt.month:02d}','聯絡人':p['contact'] or ''}
 vals.update(extra or {})
 for k,v in vals.items(): text=text.replace('{{'+k+'}}',str(v))
 return text

@app.route('/')
def index():
 q=request.args.get('q',''); return render_template('index.html',projects=all_projects(),templates=all_templates(),docs=all_docs(q),q=q)
@app.route('/projects/new',methods=['GET','POST'])
def project_new():
 if request.method=='POST':
  f=request.form;c=conn();cur=c.execute('INSERT INTO projects(name,office,year,prefix,default_recipient,speed,security,address,phone,fax,email,contact,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(f['name'],f['office'],int(f['year']),f['prefix'],f.get('default_recipient',''),f.get('speed','普通件'),f.get('security','普通'),f.get('address',''),f.get('phone',''),f.get('fax',''),f.get('email',''),f.get('contact',''),datetime.now().isoformat(timespec='seconds')));c.commit();pid=cur.lastrowid;c.close();flash('專案已建立');return redirect(url_for('project_view',pid=pid))
 return render_template('project_form.html',project=None)
@app.route('/projects/<int:pid>')
def project_view(pid):
 p=get_project(pid)
 if not p: abort(404)
 return render_template('project_view.html',project=p,templates=all_templates(),docs=project_docs(pid))
@app.route('/projects/<int:pid>/edit',methods=['GET','POST'])
def project_edit(pid):
 p=get_project(pid)
 if not p:abort(404)
 if request.method=='POST':
  f=request.form;c=conn();c.execute('UPDATE projects SET name=?,office=?,year=?,prefix=?,default_recipient=?,speed=?,security=?,address=?,phone=?,fax=?,email=?,contact=? WHERE id=?',(f['name'],f['office'],int(f['year']),f['prefix'],f.get('default_recipient',''),f.get('speed','普通件'),f.get('security','普通'),f.get('address',''),f.get('phone',''),f.get('fax',''),f.get('email',''),f.get('contact',''),pid));c.commit();c.close();flash('專案已更新');return redirect(url_for('project_view',pid=pid))
 return render_template('project_form.html',project=p)
@app.post('/projects/<int:pid>/delete')
def project_delete(pid):
 c=conn();c.execute('DELETE FROM projects WHERE id=?',(pid,));c.commit();c.close();flash('專案及其公文已刪除');return redirect(url_for('index'))

@app.route('/templates/new',methods=['GET','POST'])
@app.route('/templates/<int:tid>/edit',methods=['GET','POST'])
def template_edit(tid=None):
 t=get_template(tid) if tid else None
 if request.method=='POST':
  f=request.form;now=datetime.now().isoformat(timespec='seconds');vals=(f['name'],f.get('subject',''),f.get('explanation',''),f.get('attachment',''),f.get('originals',''),f.get('copies',''),1 if f.get('quick_enabled') else 0,f.get('quick_buttons',''),now)
  c=conn()
  if tid:c.execute('UPDATE templates SET name=?,subject=?,explanation=?,attachment=?,originals=?,copies=?,quick_enabled=?,quick_buttons=?,updated_at=? WHERE id=?',vals+(tid,))
  else:c.execute('INSERT INTO templates(name,subject,explanation,attachment,originals,copies,quick_enabled,quick_buttons,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)',vals+(now,))
  c.commit();c.close();flash('範本已儲存');return redirect(url_for('index'))
 return render_template('template_form.html',template=t)
@app.post('/templates/<int:tid>/delete')
def template_delete(tid):
 c=conn();c.execute('DELETE FROM templates WHERE id=?',(tid,));c.commit();c.close();flash('範本已刪除');return redirect(url_for('index'))

@app.route('/documents/new')
def document_new():
 ps=all_projects();pid=request.args.get('project_id',type=int) or (ps[0]['id'] if ps else None);p=get_project(pid) if pid else None
 if not p:return redirect(url_for('project_new'))
 d=date.today().isoformat();serial=next_serial(pid,d);tid=request.args.get('template_id',type=int);t=get_template(tid) if tid else None
 pre={}
 if t: pre={'subject':substitute(t['subject'],p),'explanation':substitute(t['explanation'],p),'attachment':substitute(t['attachment'],p),'originals':substitute(t['originals'],p),'copies':substitute(t['copies'],p)}
 return render_template('document_form.html',project=p,projects=ps,doc=None,serial=serial,doc_date=d,template=t,pre=pre)
@app.route('/documents/<int:did>/edit',methods=['GET','POST'])
def document_edit(did):
 c=conn();d=c.execute('SELECT * FROM documents WHERE id=?',(did,)).fetchone();c.close()
 if not d:abort(404)
 if request.method=='POST':return save_document(did)
 return render_template('document_form.html',project=get_project(d['project_id']),projects=all_projects(),doc=d,serial=d['serial'],doc_date=d['doc_date'],template=None,pre={})
def save_document(did=None):
 f=request.form;pid=int(f['project_id']);d=f['doc_date'];serial=f.get('serial') or next_serial(pid,d);now=datetime.now().isoformat(timespec='seconds');vals=(pid,d,serial,f.get('recipient',''),f.get('speed','普通件'),f.get('security','普通'),f.get('attachment',''),f.get('subject',''),f.get('explanation',''),f.get('originals',''),f.get('copies',''))
 c=conn()
 if did:c.execute('UPDATE documents SET project_id=?,doc_date=?,serial=?,recipient=?,speed=?,security=?,attachment=?,subject=?,explanation=?,originals=?,copies=?,updated_at=? WHERE id=?',vals+(now,did));new_id=did
 else:cur=c.execute('INSERT INTO documents(project_id,doc_date,serial,recipient,speed,security,attachment,subject,explanation,originals,copies,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',vals+(now,now));new_id=cur.lastrowid
 c.commit();c.close();flash('公文已儲存');return redirect(url_for('document_view',did=new_id))
@app.post('/documents/save')
def document_save():return save_document()
@app.route('/documents/<int:did>')
def document_view(did):
 c=conn();d=c.execute('SELECT d.*,p.name project_name,p.office,p.prefix,p.year,p.address,p.phone,p.fax,p.email,p.contact FROM documents d JOIN projects p ON p.id=d.project_id WHERE d.id=?',(did,)).fetchone();c.close()
 if not d:abort(404)
 return render_template('document_view.html',doc=d,no=doc_no(d,d['doc_date'],d['serial']))
@app.get('/documents/<int:did>/clone')
def document_clone(did):
 c=conn();d=c.execute('SELECT * FROM documents WHERE id=?',(did,)).fetchone();c.close()
 if not d:abort(404)
 p=get_project(d['project_id']);today=date.today().isoformat();return render_template('document_form.html',project=p,projects=all_projects(),doc=None,serial=next_serial(p['id'],today),doc_date=today,template=None,pre=dict(d))
@app.post('/documents/<int:did>/delete')
def document_delete(did):
 c=conn();d=c.execute('SELECT project_id FROM documents WHERE id=?',(did,)).fetchone();c.execute('DELETE FROM documents WHERE id=?',(did,));c.commit();c.close();flash('公文已刪除');return redirect(url_for('project_view',pid=d['project_id']) if d else url_for('index'))

@app.get('/documents/<int:did>/word')
def word(did):
 from docx.oxml import OxmlElement, parse_xml
 from docx.oxml.ns import qn
 c=conn();d=c.execute('SELECT d.*,p.name project_name,p.office,p.prefix,p.year,p.address,p.phone,p.fax,p.email,p.contact FROM documents d JOIN projects p ON p.id=d.project_id WHERE d.id=?',(did,)).fetchone();c.close()
 if not d:abort(404)
 doc=Document();sec=doc.sections[0];sec.top_margin=Mm(20);sec.bottom_margin=Mm(20);sec.left_margin=Mm(25);sec.right_margin=Mm(20)
 normal=doc.styles['Normal'];normal.font.name='標楷體';normal._element.rPr.rFonts.set(qn('w:eastAsia'),'標楷體');normal.font.size=Pt(14);normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1.35
 def sr(run,size=14,bold=False):run.font.name='標楷體';run._element.rPr.rFonts.set(qn('w:eastAsia'),'標楷體');run.font.size=Pt(size);run.bold=bold
 def para(text='',size=14,bold=False,align=None,before=0,after=0):
  p=doc.add_paragraph();p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
  if align is not None:p.alignment=align
  sr(p.add_run(text),size,bold);return p
 def field(label,value,size=12):
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.25;sr(p.add_run(label),size);sr(p.add_run(str(value or '')),size);return p
 # 裝訂線：放在頁首的絕對定位 VML，不佔正文流
 hp=sec.header.paragraphs[0];hp.paragraph_format.space_after=Pt(0)
 vml='''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape id="BindingText" style="position:absolute;margin-left:-58pt;margin-top:165pt;width:18pt;height:150pt;z-index:1;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" stroked="f" filled="f"><v:textbox inset="0,0,0,0"><w:txbxContent><w:p><w:r><w:rPr><w:rFonts w:eastAsia="標楷體"/><w:sz w:val="24"/></w:rPr><w:t>裝</w:t><w:br/><w:br/><w:t>訂</w:t><w:br/><w:br/><w:t>線</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape><v:line from="-48pt,120pt" to="-48pt,155pt" style="position:absolute;z-index:1;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" strokecolor="#777777" strokeweight="0.5pt"/><v:line from="-48pt,330pt" to="-48pt,430pt" style="position:absolute;z-index:1;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" strokecolor="#777777" strokeweight="0.5pt"/></w:pict></w:r>'''
 try:hp._p.append(parse_xml(vml))
 except Exception:pass
 para(f'{d["office"]}　函',20,True,WD_ALIGN_PARAGRAPH.CENTER,after=8)
 t=doc.add_table(rows=1,cols=2);t.alignment=WD_TABLE_ALIGNMENT.RIGHT;t.autofit=False;t.columns[0].width=Mm(68);t.columns[1].width=Mm(92)
 for cell in t.rows[0].cells:
  tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders');tcPr.append(borders)
 right=t.cell(0,1);right.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP;right.text=''
 for i,x in enumerate([f'聯絡住址：{d["address"] or ""}',f'電話：{d["phone"] or ""}　傳真：{d["fax"] or ""}',f'E-Mail：{d["email"] or ""}',f'連絡人：{d["contact"] or ""}']):
  pp=right.paragraphs[0] if i==0 else right.add_paragraph();pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.05;sr(pp.add_run(x),10)
 para('',before=6,after=2);field('受文者：',d['recipient'],16)
 dt=datetime.strptime(d['doc_date'],'%Y-%m-%d');field('發文日期：',f'中華民國 {dt.year-1911} 年 {dt.month:02d} 月 {dt.day:02d} 日');field('發文字號：',doc_no(d,d['doc_date'],d['serial']));field('速別：',d['speed']);field('密等及解密條件：',d['security']);field('附件：',d['attachment'])
 p=doc.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.35;sr(p.add_run('主旨：'),14,False);sr(p.add_run(d['subject'] or ''),14,False)
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);sr(p.add_run('說明：'),14,False)
 for x in [x.strip() for x in (d['explanation'] or '').splitlines() if x.strip()]:
  pp=doc.add_paragraph();pp.paragraph_format.left_indent=Mm(8);pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.35;sr(pp.add_run(x),14,False)
 para('',before=4,after=2);field('正本：',d['originals']);field('副本：',d['copies'])
 bio=io.BytesIO();doc.save(bio);bio.seek(0);return send_file(bio,as_attachment=True,download_name=f'{doc_no(d,d["doc_date"],d["serial"]).replace(" ","")}-公文.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

with app.app_context():init_db()
if __name__=='__main__':app.run(host='0.0.0.0',port=int(os.getenv('PORT',5000)),debug=os.getenv('FLASK_DEBUG')=='1')
